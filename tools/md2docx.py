r"""Сборка .docx из Markdown-документа «Изменения по замечаниям».

Запуск:
  set PYTHONPATH=C:\Users\shame\AppData\Local\Temp\opencode\pylibs
  .venv\Scripts\python.exe tools\md2docx.py

Как это работает:
  1. Markdown превращается в HTML (библиотека markdown).
  2. HTML разбирается в дерево и переносится в .docx (библиотека python-docx).

Зависимости (markdown, python-docx) лежат во временной папке, чтобы
не менять .venv проекта. Word для этой сборки не нужен.

Результат: ИЗМЕНЕНИЯ_ПО_ЗАМЕЧАНИЯМ.docx рядом с исходным .md.
"""

from __future__ import annotations

import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import markdown
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
SOURCE = BASE_DIR / "ИЗМЕНЕНИЯ_ПО_ЗАМЕЧАНИЯМ.md"
TARGET = BASE_DIR / "ИЗМЕНЕНИЯ_ПО_ЗАМЕЧАНИЯМ.docx"

BODY_FONT = "Segoe UI"
MONO_FONT = "Consolas"
INK = RGBColor(0x1C, 0x24, 0x30)
ACCENT = RGBColor(0x1F, 0x57, 0xC4)
MUTED = RGBColor(0x41, 0x4C, 0x5C)


class Node:
    __slots__ = ("tag", "attrs", "children")

    def __init__(self, tag: str, attrs: dict | None = None) -> None:
        self.tag = tag
        self.attrs = attrs or {}
        self.children: list = []

    def iter(self):
        yield self
        for child in self.children:
            if isinstance(child, Node):
                yield from child.iter()


VOID = {"br", "hr", "img", "meta", "link", "input"}


class TreeBuilder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, dict(attrs))
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag, dict(attrs)))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if data.strip():
            self.stack[-1].children.append(data)


def parse_html(html: str) -> Node:
    builder = TreeBuilder()
    builder.feed(html)
    builder.close()
    return builder.root


def text_of(node: Node) -> str:
    parts = []
    for child in node.children:
        if isinstance(child, str):
            parts.append(child)
        else:
            parts.append(text_of(child))
    return "".join(parts)


def norm_text(text: str) -> str:
    """Схлопывает переносы строк и отступы markdown в пробелы.

    Пробелы по краям сохраняются, чтобы слова не слипались при склейке
    соседних кусков (например, текст + <code>).
    """
    if not text:
        return ""
    if not text.strip():
        return ""
    lead = " " if text[0].isspace() else ""
    trail = " " if text[-1].isspace() else ""
    return lead + re.sub(r"\s+", " ", text.strip()) + trail


def inline_runs(node: Node, bold=False, italic=False, mono=False) -> list[tuple[str, bool, bool, bool]]:
    """Возвращает список кортежей (текст, жирный, курсив, моноширинный)."""
    runs: list[tuple[str, bool, bool, bool]] = []
    for child in node.children:
        if isinstance(child, str):
            runs.append((norm_text(child), bold, italic, mono))
        elif child.tag in ("strong", "b"):
            runs.extend(inline_runs(child, True, italic, mono))
        elif child.tag in ("em", "i"):
            runs.extend(inline_runs(child, bold, True, mono))
        elif child.tag == "code":
            runs.extend(inline_runs(child, bold, italic, True))
        elif child.tag == "br":
            runs.append(("\n", bold, italic, mono))
        else:
            runs.extend(inline_runs(child, bold, italic, mono))
    return runs


def add_runs(paragraph, runs, size=None, color=None):
    for text, bold, italic, mono in runs:
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x4B)
        else:
            run.font.name = BODY_FONT
            if size:
                run.font.size = size
            if color:
                run.font.color.rgb = color
    return paragraph


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before in (
        ("Heading 1", 20, INK, 0),
        ("Heading 2", 15, ACCENT, 16),
        ("Heading 3", 12.5, INK, 12),
    ):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def add_paragraph_border(paragraph, color="D7DDE6", size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = borders.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def shade_paragraph(paragraph, fill="F2F4F8"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def render_table(doc: Document, node: Node) -> None:
    header_cells: list[Node] = []
    body_rows: list[list[Node]] = []
    for child in node.children:
        if not isinstance(child, Node):
            continue
        if child.tag == "thead":
            for tr in child.children:
                if isinstance(tr, Node) and tr.tag == "tr":
                    header_cells = [td for td in tr.children if isinstance(td, Node)]
        elif child.tag == "tbody":
            for tr in child.children:
                if isinstance(tr, Node) and tr.tag == "tr":
                    body_rows.append([td for td in tr.children if isinstance(td, Node)])

    if not body_rows and not header_cells:
        return

    cols = max([len(header_cells)] + [len(r) for r in body_rows] or [1])
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def fill(cells, nodes, bold: bool) -> None:
        for i in range(cols):
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            if i < len(nodes):
                runs = inline_runs(nodes[i], bold=bold)
                for text, b, it, mono in runs:
                    run = para.add_run(text)
                    run.bold = b
                    run.italic = it
                    run.font.name = MONO_FONT if mono else BODY_FONT
                    run.font.size = Pt(9.5 if mono else 10)
                    run.font.color.rgb = RGBColor(0xA3, 0x1D, 0x4B) if mono else INK

    if header_cells:
        fill(table.add_row().cells, header_cells, bold=True)
        for cell in table.rows[0].cells:
            shade_paragraph(cell.paragraphs[0], "EEF3FB")

    for row in body_rows:
        fill(table.add_row().cells, row, bold=False)

    doc.add_paragraph()


def render_list(doc: Document, node: Node, ordered: bool, depth: int = 0) -> None:
    style = "List Number" if ordered else "List Bullet"
    for li in node.children:
        if not isinstance(li, Node) or li.tag != "li":
            continue
        para = doc.add_paragraph(style=style)
        if depth:
            para.paragraph_format.left_indent = Cm(0.75 * depth)
        for child in li.children:
            if isinstance(child, Node) and child.tag in ("ul", "ol"):
                render_list(doc, child, child.tag == "ol", depth + 1)
                continue
            if isinstance(child, Node) and child.tag == "p":
                add_runs(para, inline_runs(child))
            elif isinstance(child, str):
                para.add_run(norm_text(child))
            elif isinstance(child, Node):
                # Плотный пункт списка: <strong>, <code> и прочее лежат прямо в <li>
                add_runs(para, inline_runs(child))


def render_block(doc: Document, node: Node) -> None:
    for child in node.children:
        if isinstance(child, str):
            text = child.strip()
            if text:
                doc.add_paragraph(text)
            continue

        tag = child.tag
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(tag[1])
            para = doc.add_heading("", level=min(level, 4))
            add_runs(para, inline_runs(child))
        elif tag == "p":
            para = doc.add_paragraph()
            add_runs(para, inline_runs(child))
        elif tag == "pre":
            code = text_of(child).strip("\n")
            for line in code.split("\n"):
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.left_indent = Cm(0.4)
                run = para.add_run(line if line else " ")
                run.font.name = MONO_FONT
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
                shade_paragraph(para, "F2F4F8")
            doc.add_paragraph()
        elif tag == "ul":
            render_list(doc, child, ordered=False)
        elif tag == "ol":
            render_list(doc, child, ordered=True)
        elif tag == "table":
            render_table(doc, child)
        elif tag == "blockquote":
            for sub in child.children:
                if isinstance(sub, Node) and sub.tag == "p":
                    para = doc.add_paragraph()
                    para.paragraph_format.left_indent = Cm(0.6)
                    add_runs(para, inline_runs(sub, italic=True), color=MUTED)
        elif tag == "hr":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(8)
            add_paragraph_border(para)
        else:
            render_block(doc, child)


def build() -> Path:
    text = SOURCE.read_text(encoding="utf-8")
    html = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="html5",
    )
    tree = parse_html(html)

    doc = Document()
    style_document(doc)
    render_block(doc, tree)

    doc.save(str(TARGET))
    return TARGET


def main() -> int:
    if not SOURCE.exists():
        print(f"нет исходника: {SOURCE}")
        return 1
    path = build()
    print(f"Сохранено: {path} ({path.stat().st_size} байт)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
